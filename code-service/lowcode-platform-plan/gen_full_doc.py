# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return h

def add_para(text, bold=False, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return p

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Cm(1)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = '微软雅黑'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r+1].cells[c]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = '微软雅黑'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return table

# ========== 封面 ==========
title = doc.add_heading('', level=0)
run = title.add_run('低代码平台\n完整开发文档')
run.font.size = Pt(26)
run.font.name = '微软雅黑'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_para('')
add_para('版本：v5.0（最终版）    日期：2026-09-22    模块：bcd-code（端口 3015）', size=11)
add_para('数据源：lowcode（元数据库）/ project（业务库）', size=11)
add_para('', size=10)
add_para('整合内容：', bold=True, size=11)
add_para('  · 元数据驱动动态 API（零代码 CRUD + 对外发布）', size=10)
add_para('  · 通用触发源（按钮点击 / 状态变更 / 字段变更 / API调用）', size=10)
add_para('  · 动作执行体系（多动作顺序执行 + 多插件混配 + 动作间传数据）', size=10)
add_para('  · 事务化设计（compileXxx + executeInTransaction 原子提交）', size=10)
add_para('  · 三层扩展体系（表达式 / Groovy脚本 / JAR插件）', size=10)
add_para('  · 审批流（从零搭建，线性审批 MVP）', size=10)
add_para('  · 微服务插件分发（MinIO + Nacos 自动广播）', size=10)
add_para('  · 缓存复用（L1Caffeine + L2Redis + BatchQueryStrategy）', size=10)

doc.add_page_break()

# ========== 目录 ==========
add_heading('目录', 1)
chapters = [
    '第一章  整体架构',
    '第二章  元数据驱动动态 API',
    '第三章  通用触发源（4种触发类型）',
    '第四章  动作执行体系（多动作 + 多插件）',
    '第五章  事务化设计',
    '第六章  三层扩展体系（表达式/脚本/插件）',
    '第七章  插件开发全流程',
    '第八章  插件能调用什么、怎么调用',
    '第九章  审批流设计',
    '第十章  缓存设计（复用现有设施）',
    '第十一章  微服务插件分发',
    '第十二章  调试方案',
    '第十三章  风险与对策',
    '第十四章  依赖清单',
    '第十五章  落地步骤',
    '第十六章  FAQ',
]
for ch in chapters:
    add_para(ch, size=11)

doc.add_page_break()

# ========== 第一章 ==========
add_heading('第一章  整体架构', 1)

add_heading('1.1 架构总览', 2)
add_code('''外部消费者
  ↓ POST /api/open/**（元数据驱动 CRUD）
  ↓ POST /api/open/trigger/{triggerId}（按钮/API触发）
OpenApiController / TriggerApiController
  ↓
MetaEngineService（元数据驱动执行引擎）
  ├── FieldGuard          字段白名单过滤
  ├── FieldTransformer     字段重命名 + code→id 转换
  ├── BatchFieldTransformer 批量转换（复用 BatchQueryStrategy）
  ├── SqlEngineFacade      编译参数化 SQL → 执行
  └── StatusChangeDetector 检测字段变更 → 触发动作
  ↓
TriggerDispatcher（触发分发，按顺序执行动作列表）
  ├── sync（同步，编译SQL加入事务）
  │   ├── updateField     改字段值
  │   ├── docConvert      单据转换
  │   ├── approval        触发审批
  │   ├── script          Groovy 脚本
  │   └── plugin          JAR 插件（可多个）
  └── async（异步，发 RocketMQ）
      ├── httpCall        调外部接口
      ├── mqSend          发消息
      └── notify          通知
  ↓
executeInTransaction（一次性原子提交）''')

add_heading('1.2 核心原则', 2)
add_table(['原则', '说明'], [
    ['平台代码不修改', '交付后所有自定义逻辑通过配置/脚本/插件实现'],
    ['元数据驱动', '表名/字段/数据源/操作类型/字段映射全部来自元数据'],
    ['复用不造轮子', '缓存复用L1Caffeine+L2Redis，查询复用BatchQueryStrategy'],
    ['同步/异步分离', '影响一致性的同步(事务内)，纯通知的异步(RocketMQ)'],
    ['按钮直接传数据', '前端按钮点击直接传当前行数据，后端不查库'],
])

add_heading('1.3 现有复用资产', 2)
add_table(['资产', '复用于'], [
    ['SqlEngineFacade', '执行引擎内核（query/insert/update/delete/batchInsert）'],
    ['SqlFieldValidator', '表/列合法性校验（5min TTL 缓存）'],
    ['SqlTemplateEngine', 'custom SQL 模板渲染（@{alias.field}@ 占位符）'],
    ['LowcodeHelper', '元数据表 CRUD（dbName=lowcode）'],
    ['L1CaffeineCache', '元数据缓存 + code→id L1'],
    ['L2RedisCache', 'code→id L2（Redisson 管道化 mget/mset）'],
    ['KeyBuilder', '缓存 key 构造 bd:<cacheName>:s<ver>:<id>'],
    ['BatchQueryStrategy', '批量 IN 查询回源（默认 1000/批）'],
    ['CacheInvalidationListener', 'RocketMQ 跨实例缓存失效广播'],
    ['ResMsg', '统一返回对象'],
    ['RocketMQTemplate', '异步动作 + 跨实例广播'],
    ['MinioService', '插件 JAR 存储'],
    ['NacosConfigService', '插件配置广播 + 各服务监听'],
])

doc.add_page_break()

# ========== 第二章 ==========
add_heading('第二章  元数据驱动动态 API', 1)

add_heading('2.1 数据库设计（DDL）', 2)
add_para('数据模型元数据表：')
add_code('''CREATE TABLE IF NOT EXISTS lowcode_model (
    model_code      VARCHAR(128) PRIMARY KEY,
    model_name      VARCHAR(255),
    table_name      VARCHAR(255) NOT NULL,
    db_name         VARCHAR(64) DEFAULT 'project',
    fields_json     TEXT,           -- 字段元数据数组
    relations_json  TEXT,           -- 关联关系（可选）
    cache_enabled   BOOLEAN DEFAULT FALSE,
    cache_name      VARCHAR(128),
    code_column     VARCHAR(128),
    id_column       VARCHAR(128),
    enabled         BOOLEAN DEFAULT TRUE,
    created_at      VARCHAR(64),
    updated_at      VARCHAR(64)
);''')

add_para('fields_json 结构（name 用表列名 snake_case）：')
add_code('''[
  {"name":"order_id","type":"string","primaryKey":true,"comment":"单据ID"},
  {"name":"warehouse_id","type":"string","comment":"仓库ID",
   "ref":{"model":"wms_warehouse","idColumn":"warehouse_id","codeColumn":"warehouse_code"}},
  {"name":"status","type":"string","comment":"状态"}
]''')

add_para('API 注册表扩展：')
add_code('''ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS api_code      VARCHAR(128);
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS model_code     VARCHAR(128);
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS operation      VARCHAR(32);
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS param_mapping  TEXT;
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS sql_template   TEXT;
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS db_name        VARCHAR(64);
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS enabled        BOOLEAN DEFAULT TRUE;
ALTER TABLE lowcode_api_registry ADD COLUMN IF NOT EXISTS max_batch_size INT DEFAULT 1000;''')

add_heading('2.2 对外 API 入口', 2)
add_table(['操作', '路径', '入参', '说明'], [
    ['查询', 'POST /api/open/{path}', '{conditions,orderBy,page,pageSize}', '分页查询'],
    ['新增', 'POST /api/open/{path}', '{data:{字段:值}}', '单条新增'],
    ['批量新增', 'POST /api/open/{path}', '{data:[...1000条]}', '批量新增'],
    ['修改', 'POST /api/open/{path}', '{data,conditions}', 'conditions非空'],
    ['删除', 'POST /api/open/{path}', '{conditions}', 'conditions非空'],
])

add_heading('2.3 字段转换（code→id）', 2)
add_para('外部传编码（如仓库编码 WH01），内部自动查表转 ID。param_mapping 配置：')
add_code('''{
  "warehouseCode": {"targetField":"warehouse_id","transform":"codeToId","refModel":"wms_warehouse"},
  "orderNo":       {"targetField":"order_no","transform":"direct"}
}''')

add_para('批量转换性能优化（复用 BatchQueryStrategy）：')
add_table(['方案', 'DB往返', '1000万行物料表表现'], [
    ['sqlTemplate 子查询（禁用）', '1000次子查询', '卡死'],
    ['BatchFieldTransformer（推荐）', '2次查询+1次批量插入', '毫秒级'],
])

doc.add_page_break()

# ========== 第三章 ==========
add_heading('第三章  通用触发源（4种触发类型）', 1)

add_heading('3.1 触发配置表', 2)
add_code('''CREATE TABLE IF NOT EXISTS lowcode_action_trigger (
    trigger_id      VARCHAR(64) PRIMARY KEY,
    trigger_name    VARCHAR(255),
    trigger_type    VARCHAR(32) NOT NULL,   -- buttonClick/statusChange/fieldChange/apiCall
    model_code      VARCHAR(128) NOT NULL,
    page_id         VARCHAR(64),            -- 关联页面（buttonClick必填）
    button_key      VARCHAR(128),           -- 按钮标识（buttonClick必填）
    field_name      VARCHAR(128),           -- 变更字段名（statusChange/fieldChange填）
    from_value      VARCHAR(255),           -- 源值（*=任意）
    to_value        VARCHAR(255),           -- 目标值
    actions_json    TEXT,                   -- 动作列表（有序）
    enabled         BOOLEAN DEFAULT TRUE,
    created_at      VARCHAR(64),
    updated_at      VARCHAR(64)
);''')

add_heading('3.2 四种触发源', 2)
add_table(['trigger_type', '触发时机', '必填字段', '入口'], [
    ['buttonClick', '用户点页面按钮', 'page_id, button_key', 'POST /api/open/trigger/{triggerId}'],
    ['statusChange', 'update后检测status字段变更', 'field_name, from_value, to_value', 'MetaEngineService.doUpdate自动检测'],
    ['fieldChange', 'update后检测任意字段变更', 'field_name, from_value, to_value', 'MetaEngineService.doUpdate自动检测'],
    ['apiCall', '外部系统调触发接口', '无额外', 'POST /api/open/trigger/{triggerId}'],
])

add_heading('3.3 按钮点击触发（前端直接传数据，不查库）', 2)
add_para('前端：按钮 props 配 triggerId，点击时把当前行数据直接传给后端：')
add_code('''// 前端按钮点击
async function onButtonClick(button) {
  const res = await apiCall('/api/open/trigger/' + button.props.triggerId, {
    body: JSON.stringify({
      bizId: currentRow.order_id,
      data: { ...currentRow }    // 前端已有的数据，直接传
    })
  });
}''')

add_para('后端：直接用前端传的数据，不查库：')
add_code('''@PostMapping("/trigger/{triggerId}")
public ResMsg fire(@PathVariable String triggerId, @RequestBody JSONObject body) {
    JSONObject trigger = metaCache.getTrigger(triggerId);
    String modelCode = trigger.getString("modelCode");
    String bizId = body.getString("bizId");
    JSONObject newData = body.getJSONObject("data");  // 直接用前端数据

    ActionContext ctx = new ActionContext();
    ctx.setModelCode(modelCode);
    ctx.setNewData(newData);
    ctx.setOldData(newData);
    ctx.setBizId(bizId);
    ctx.setDbName(resolveDbName(modelCode));
    ctx.setSqlList(new ArrayList<>());
    ctx.setTriggerChain(new ArrayList<>());

    dispatcher.dispatch(trigger, ctx);  // 执行动作列表

    if (!ctx.getSqlList().isEmpty()) {
        platformApi.executeInTransaction(ctx.getDbName(), ctx.getSqlList());
    }
    // ...
}''')

doc.add_page_break()

# ========== 第四章 ==========
add_heading('第四章  动作执行体系（多动作 + 多插件）', 1)

add_heading('4.1 动作列表按顺序执行', 2)
add_para('一个按钮的触发配置里，actions_json 是数组，可以配任意多个动作，多个插件混着配：')
add_code('''{
  "actionsJson": [
    {"actionType":"updateField", "name":"状态改为已确认", "execMode":"sync",
     "config":{"field":"status","value":"confirmed"}},
    {"actionType":"docConvert",  "name":"转出库单",     "execMode":"sync",
     "config":{"targetModel":"wms_outbound_order","fieldMapping":{...}}},
    {"actionType":"stockDeduct", "name":"扣减库存",     "execMode":"sync","config":{}},
    {"actionType":"stockLog",    "name":"写库存日志",   "execMode":"sync","config":{}},
    {"actionType":"approval",    "name":"提交审批",     "execMode":"sync",
     "config":{"flowCode":"inbound_approval","onApprove":"approved","onReject":"draft"}},
    {"actionType":"httpCall",    "name":"通知WCS",     "execMode":"async",
     "config":{"url":"http://wcs/api/notify","method":"POST"}}
  ]
}''')

add_table(['动作', '类型', '来源', '执行方式'], [
    ['状态改为已确认', 'updateField', '平台内置', '同步'],
    ['转出库单', 'docConvert', '平台内置', '同步'],
    ['扣减库存', 'stockDeduct', '插件1', '同步'],
    ['写库存日志', 'stockLog', '插件2', '同步'],
    ['提交审批', 'approval', '平台内置', '同步'],
    ['通知WCS', 'httpCall', '平台内置', '异步(MQ)'],
])

add_para('同步动作（含多个插件）全部在同一事务里，按顺序执行。异步动作事务提交后发MQ。')

add_heading('4.2 动作间数据传递（ctx.variables）', 2)
add_para('每个动作执行后可把结果放入 ctx.variables，下一个动作从中取用：')
add_code('''// 动作1：单据转换，把结果存入 variables
public class DocConvertActionExecutor implements ActionExecutor {
    public Object execute(JSONObject config, ActionContext ctx, PlatformApi api) {
        // ... 编译 insert SQL 加入事务 ...
        ctx.getVariables().put("convertedOrderId", newOrderId);
        ctx.getVariables().put("convertedOrderData", targetData);
        return null;
    }
}

// 动作2：你的插件，从 variables 拿转换结果
public class StockDeductExecutor implements ActionExecutor {
    public Object execute(JSONObject config, ActionContext ctx, PlatformApi api) {
        String convertedOrderId = (String) ctx.getVariables().get("convertedOrderId");
        JSONObject convertedOrder = (JSONObject) ctx.getVariables().get("convertedOrderData");
        // 用转换后的数据执行你的逻辑...
    }
}''')

add_heading('4.3 TriggerDispatcher 分发逻辑', 2)
add_code('''public void dispatch(JSONObject trigger, ActionContext ctx) {
    JSONArray actions = JSON.parseArray(trigger.getString("actionsJson"));
    for (int i = 0; i < actions.size(); i++) {
        JSONObject action = actions.getJSONObject(i);
        String type = action.getString("actionType");
        String execMode = action.getString("execMode");

        // 查找执行器：先查Spring注入的，再查插件加载的
        ActionExecutor executor = executors.get(type);
        if (executor == null) executor = pluginManager.getExecutor(type);

        if ("async".equals(execMode)) {
            sendAsync(action, ctx);           // 发 RocketMQ
        } else {
            executor.execute(action.getJSONObject("config"), ctx, platformApi);
            // 同步动作编译SQL加入ctx.sqlList，不直接执行
        }
    }
}''')

doc.add_page_break()

# ========== 第五章 ==========
add_heading('第五章  事务化设计', 1)

add_heading('5.1 原理', 2)
add_para('写操作不直接执行，编译成 CompiledSql 收集到 sqlList，最后 executeInTransaction 一次性提交。读操作直接执行（不进事务）。')

add_heading('5.2 PlatformApi 接口', 2)
add_code('''public interface PlatformApi {
    // 读操作：直接执行（不进事务）
    JSONArray query(JSONObject req, String dbName) throws Exception;
    long count(JSONObject req, String dbName) throws Exception;

    // 写操作编译模式：返回 CompiledSql，不执行（加入事务）
    CompiledSql compileInsert(JSONObject req, String dbName) throws Exception;
    CompiledSql compileUpdate(JSONObject req, String dbName) throws Exception;
    CompiledSql compileDelete(JSONObject req, String dbName) throws Exception;

    // 事务执行
    void executeInTransaction(String dbName, List<CompiledSql> sqlList) throws Exception;

    // 其他
    String redisGet(String key);
    void redisSet(String key, String value, long ttlSec);
    void sendMq(String topic, Object message);
    JSONObject getModel(String modelCode) throws Exception;
    Object executeApi(String apiCode, JSONObject body) throws Exception;
}''')

add_heading('5.3 事务模式核心规则', 2)
add_table(['场景', '用什么', '为什么'], [
    ['查数据', 'api.query()', '读操作不需事务保护，直接执行最快'],
    ['增删改（同步动作）', 'api.compileXxx() + ctx.addSql()', '编译加入事务，平台统一executeInTransaction，保证原子性'],
    ['发消息/调外部接口', '异步动作里用 api.sendMq()', '事务提交后才发，避免回滚但消息已发'],
    ['缓存读写', 'api.redisGet/Set()', '缓存不进事务，独立操作'],
])

doc.add_page_break()

# ========== 第六章 ==========
add_heading('第六章  三层扩展体系（表达式/脚本/插件）', 1)

add_table(['层级', '引擎', '适用', '前端形态', '谁来写'], [
    ['表达式', 'Aviator', '字段计算、条件判断', '表达式输入框', '配置人员'],
    ['脚本', 'Groovy', '中等复杂逻辑', 'Monaco代码编辑器', '配置人员/开发'],
    ['插件', 'PF4J JAR', '重量级业务扩展', '上传JAR管理页', '二次开发者'],
])

add_heading('6.1 表达式引擎（Aviator）', 2)
add_code('''// 前端配置
{"field":"totalAmount","expression":"qty * unitPrice"}
{"condition":"status == 'draft' && totalAmount > 1000"}

// 后端执行
Expression exp = AviatorEvaluator.compile("qty * unitPrice");
Object result = exp.execute(env);''')

add_heading('6.2 脚本引擎（Groovy）', 2)
add_code('''// 前端配置（存元数据）
{"actionType":"script","config":{"lang":"groovy","script":
  "def stocks = api.query(req, ctx.dbName)\\nfor (s in stocks) { ... }"}}

// ScriptActionExecutor：编译缓存 + 注入平台能力
@Component
public class ScriptActionExecutor implements ActionExecutor {
    private ConcurrentHashMap<String, Class<?>> scriptCache = new ConcurrentHashMap<>();

    public Object execute(JSONObject config, ActionContext ctx, PlatformApi api) {
        String script = config.getString("script");
        String cacheKey = DigestUtils.md5Hex(script);
        Class<?> scriptClass = scriptCache.get(cacheKey);
        if (scriptClass == null) {
            GroovyClassLoader gcl = new GroovyClassLoader(getClass().getClassLoader());
            scriptClass = gcl.parseClass(script);
            scriptCache.put(cacheKey, scriptClass);
        }
        Script instance = (Script) scriptClass.newInstance();
        Binding binding = new Binding();
        binding.setVariable("ctx", ctx);
        binding.setVariable("api", api);
        instance.setBinding(binding);
        return instance.run();
    }
}''')

add_heading('6.3 插件机制（PF4J JAR）', 2)
add_para('插件是独立 Maven 项目，只依赖平台接口 JAR，打 JAR 上传，热加载：')
add_code('''// 插件实现 ActionExecutor 接口
public class StockAdjustExecutor implements ActionExecutor {
    public String getType() { return "stockAdjust"; }
    public Object execute(JSONObject config, ActionContext ctx, PlatformApi api) {
        // 通过 api 访问平台能力，不依赖 Spring
        JSONArray stocks = api.query(req, ctx.getDbName());
        // ...
    }
}

// SPI 注册：META-INF/services/com.api.lowcode.trigger.action.ActionExecutor
// 内容：com.customer.plugin.StockAdjustExecutor

// 打 JAR → 放 plugins/ 目录 → 自动加载''')

doc.add_page_break()

# ========== 第七章 ==========
add_heading('第七章  插件开发全流程', 1)

add_heading('7.1 环境准备', 2)
add_table(['工具', '版本', '说明'], [
    ['JDK', '8', '与平台一致'],
    ['Maven', '3.6+', '打包用'],
    ['IDE', 'IDEA/Eclipse', '开发调试'],
    ['plugin-api JAR', '1.0', '平台提供，含接口定义'],
])

add_heading('7.2 创建项目', 2)
add_code('''my-stock-plugin/
├── pom.xml
└── src/main/java/com/customer/plugin/
    └── StockDeductExecutor.java
└── src/main/resources/META-INF/services/
    └── com.api.lowcode.trigger.action.ActionExecutor''')

add_heading('7.3 pom.xml', 2)
add_code('''<project>
    <groupId>com.customer</groupId>
    <artifactId>my-stock-plugin</artifactId>
    <version>1.0.0</version>
    <packaging>jar</packaging>
    <properties>
        <maven.compiler.source>1.8</maven.compiler.source>
        <maven.compiler.target>1.8</maven.compiler.target>
    </properties>
    <dependencies>
        <dependency>
            <groupId>com.api</groupId>
            <artifactId>plugin-api</artifactId>
            <version>1.0</version>
            <scope>provided</scope>
        </dependency>
        <dependency>
            <groupId>com.alibaba.fastjson2</groupId>
            <artifactId>fastjson2</artifactId>
            <version>2.0.47</version>
            <scope>provided</scope>
        </dependency>
    </dependencies>
</project>''')

add_heading('7.4 完整示例：库存扣减插件', 2)
add_code('''package com.customer.plugin;

import com.alibaba.fastjson2.JSONArray;
import com.alibaba.fastjson2.JSONObject;
import com.api.lowcode.trigger.action.ActionExecutor;
import com.api.lowcode.trigger.action.ActionContext;
import com.api.lowcode.trigger.action.PlatformApi;
import com.server.sqlengine.model.CompiledSql;

public class StockDeductExecutor implements ActionExecutor {

    @Override
    public String getType() { return "stockDeduct"; }

    @Override
    public Object execute(JSONObject config, ActionContext ctx, PlatformApi api) throws Exception {
        JSONObject newData = ctx.getNewData();
        String warehouseId = newData.getString("warehouse_id");
        String orderId = newData.getString("order_id");
        String dbName = ctx.getDbName();

        // 1. 查入库单明细（读操作，直接执行）
        JSONObject detailReq = new JSONObject();
        detailReq.put("tableName", "wms_inbound_order_detail");
        detailReq.put("columns", JSONArray.of("material_id", "qty"));
        detailReq.put("conditions", buildEq("order_id", orderId));
        JSONArray details = api.query(detailReq, dbName);

        // 2. 逐条扣减库存（写操作，编译加入事务）
        for (int i = 0; i < details.size(); i++) {
            JSONObject detail = details.getJSONObject(i);
            String materialId = detail.getString("MATERIAL_ID");
            int qty = detail.getIntValue("QTY");

            // 查当前库存
            JSONObject stockReq = new JSONObject();
            stockReq.put("tableName", "wms_stock");
            stockReq.put("columns", JSONArray.of("stock_id", "available_qty"));
            stockReq.put("conditions", buildEq2("warehouse_id", warehouseId, "material_id", materialId));
            JSONArray stocks = api.query(stockReq, dbName);
            if (stocks.isEmpty()) throw new RuntimeException("库存不存在");

            JSONObject stock = stocks.getJSONObject(0);
            int currentQty = stock.getIntValue("AVAILABLE_QTY");
            if (currentQty < qty) throw new RuntimeException("库存不足");

            // 编译扣减SQL（加入事务，不直接执行）
            JSONObject updateReq = new JSONObject();
            updateReq.put("tableName", "wms_stock");
            updateReq.put("data", JSONObject.of("available_qty", currentQty - qty));
            updateReq.put("conditions", buildEq("stock_id", stock.getString("STOCK_ID")));
            ctx.addSql(api.compileUpdate(updateReq, dbName));

            // 编译变动记录SQL（加入事务）
            JSONObject logReq = new JSONObject();
            logReq.put("tableName", "wms_stock_log");
            logReq.put("data", JSONObject.of("stock_id", stock.getString("STOCK_ID"),
                "order_id", orderId, "change_qty", -qty, "type", "inbound_deduct"));
            ctx.addSql(api.compileInsert(logReq, dbName));
        }
        return JSONObject.of("processed", details.size());
    }

    private JSONArray buildEq(String col, Object val) {
        JSONArray c = new JSONArray();
        c.add(JSONObject.of("Symbol", "and"));
        c.add(JSONObject.of("Id", col, "Symbol", "=", "Val", val, "TableAlias", "T"));
        return c;
    }
    private JSONArray buildEq2(String c1, Object v1, String c2, Object v2) {
        JSONArray c = new JSONArray();
        c.add(JSONObject.of("Symbol", "and"));
        c.add(JSONObject.of("Id", c1, "Symbol", "=", "Val", v1, "TableAlias", "T"));
        c.add(JSONObject.of("Symbol", "and"));
        c.add(JSONObject.of("Id", c2, "Symbol", "=", "Val", v2, "TableAlias", "T"));
        return c;
    }
}''')

add_heading('7.5 SPI 注册', 2)
add_code('''# 文件：src/main/resources/META-INF/services/com.api.lowcode.trigger.action.ActionExecutor
com.customer.plugin.StockDeductExecutor''')

add_heading('7.6 打包部署', 2)
add_code('''mvn clean package
# 产出：target/my-stock-plugin-1.0.0.jar
# 上传到平台设计器插件管理页，或放 plugins/ 目录''')

doc.add_page_break()

# ========== 第八章 ==========
add_heading('第八章  插件能调用什么、怎么调用', 1)

add_table(['方法', '类型', '说明', '是否进事务'], [
    ['api.query(req, dbName)', '读', '查询数据库，返回JSONArray', '否，直接执行'],
    ['api.count(req, dbName)', '读', 'COUNT查询，返回long', '否'],
    ['api.compileInsert(req, dbName)', '写', '编译INSERT，返回CompiledSql', '是，需ctx.addSql()'],
    ['api.compileUpdate(req, dbName)', '写', '编译UPDATE，返回CompiledSql', '是，需ctx.addSql()'],
    ['api.compileDelete(req, dbName)', '写', '编译DELETE，返回CompiledSql', '是，需ctx.addSql()'],
    ['api.executeInTransaction(dbName, sqlList)', '事务', '一次性提交所有SQL', '平台统一调用'],
    ['api.redisGet(key)', '缓存', '读Redis', '否'],
    ['api.redisSet(key, value, ttl)', '缓存', '写Redis', '否'],
    ['api.sendMq(topic, message)', '消息', '发RocketMQ', '否（异步动作用）'],
    ['api.getModel(modelCode)', '元数据', '读模型元数据（走缓存）', '否'],
    ['api.executeApi(apiCode, body)', '引擎', '调元数据驱动API', '否'],
])

add_heading('8.1 查询示例', 2)
add_code('''JSONObject req = new JSONObject();
req.put("tableName", "wms_stock");
req.put("columns", JSONArray.of("stock_id", "available_qty"));
req.put("conditions", buildEq("warehouse_id", "wh001"));
req.put("page", 1);
req.put("pageSize", 100);
JSONArray result = api.query(req, "project");''')

add_heading('8.2 写操作示例（编译加入事务）', 2)
add_code('''JSONObject req = new JSONObject();
req.put("tableName", "wms_stock");
req.put("data", JSONObject.of("available_qty", 99));
req.put("conditions", buildEq("stock_id", "stk001"));
CompiledSql sql = api.compileUpdate(req, ctx.getDbName());
ctx.addSql(sql);  // 加入事务，由平台统一executeInTransaction''')

doc.add_page_break()

# ========== 第九章 ==========
add_heading('第九章  审批流设计', 1)

add_heading('9.1 审批流表（4张表，从零搭建）', 2)
add_code('''CREATE TABLE IF NOT EXISTS lowcode_approval_flow (
    flow_code VARCHAR(64) PRIMARY KEY, flow_name VARCHAR(255),
    model_code VARCHAR(128), enabled BOOLEAN DEFAULT TRUE, created_at VARCHAR(64)
);
CREATE TABLE IF NOT EXISTS lowcode_approval_node (
    node_id VARCHAR(64) PRIMARY KEY, flow_code VARCHAR(64) NOT NULL,
    node_name VARCHAR(255), node_order INT NOT NULL,
    approver_type VARCHAR(32), approver_value VARCHAR(512), created_at VARCHAR(64)
);
CREATE TABLE IF NOT EXISTS lowcode_approval_instance (
    instance_id VARCHAR(64) PRIMARY KEY, flow_code VARCHAR(64) NOT NULL,
    model_code VARCHAR(128) NOT NULL, biz_id VARCHAR(64) NOT NULL,
    current_node_id VARCHAR(64), status VARCHAR(32) DEFAULT 'pending',
    created_at VARCHAR(64), updated_at VARCHAR(64)
);
CREATE TABLE IF NOT EXISTS lowcode_approval_record (
    record_id VARCHAR(64) PRIMARY KEY, instance_id VARCHAR(64) NOT NULL,
    node_id VARCHAR(64) NOT NULL, approver VARCHAR(64) NOT NULL,
    action VARCHAR(32) NOT NULL, comment TEXT, created_at VARCHAR(64)
);''')

add_heading('9.2 审批流程', 2)
add_code('''状态变更 → approval动作 → 创建审批实例，单据状态改"approving"
  ↓
审批人操作 POST /api/lowcode/approval/approve
  ↓
全部节点通过 → 单据状态→onApprove(如"approved") → 可再触发新动作
任一节点拒绝 → 单据状态→onReject(如"draft")  → 可再触发新动作''')

add_heading('9.3 MVP范围', 2)
add_table(['做', '不做（后续）'], [
    ['线性审批（按节点顺序）', '会签/分支审批'],
    ['user类型审批人（指定用户ID）', 'role/dynamic审批人解析'],
    ['审批通过/拒绝触发新状态变更', '审批委托/转签'],
])

doc.add_page_break()

# ========== 第十章 ==========
add_heading('第十章  缓存设计（复用现有设施）', 1)

add_heading('10.1 元数据缓存（复用 L1CaffeineCache）', 2)
add_para('缓存 lowcode_model 和 lowcode_api_registry 查询结果，TTL 5分钟，变更时 evict + RocketMQ 跨实例广播。')

add_heading('10.2 code→id 缓存（复用 L1+L2+BatchQueryStrategy）', 2)
add_table(['层级', '设施', '说明'], [
    ['L1', 'L1CaffeineCache', '本地缓存，命中直接返回'],
    ['L2', 'L2RedisCache', 'Redis管道化mget/mset，跨实例共享'],
    ['DB', 'BatchQueryStrategy', 'IN分批查询回源（默认1000/批）'],
])

add_heading('10.3 分层缓存策略', 2)
add_table(['基础数据', '数据量', '策略'], [
    ['仓库/单位', '< 1万', 'L1全量缓存（cache_enabled=true）'],
    ['物料/客户', '10万~1000万', '批量IN查询，不缓存全量，结果回填L2'],
    ['超大数据', '> 1000万', '批量IN + Redis热点缓存'],
])

doc.add_page_break()

# ========== 第十一章 ==========
add_heading('第十一章  微服务插件分发', 1)

add_heading('11.1 分发架构', 2)
add_code('''上传JAR → MinIO + lowcode_plugin表 + Nacos配置
  ↓ Nacos广播
各服务监听 → 按target_services过滤 → checksum对比 → MinIO下载 → 热加载''')

add_heading('11.2 Nacos配置（dataId: lowcode.plugins）', 2)
add_code('''{
  "plugins": [{
    "pluginId": "stock-adjust",
    "version": "1.0.0",
    "jarUrl": "minio://lowcode/plugins/stock-adjust-1.0.0.jar",
    "checksum": "abc123",
    "targetServices": ["bcd-code", "bcd-wms"],
    "enabled": true
  }]
}''')

add_heading('11.3 各服务监听热加载', 2)
add_code('''@NacosConfigListener(dataId = "lowcode.plugins", groupId = "LOWCODE")
public void onPluginConfigChanged(String config) {
    JSONArray plugins = JSON.parseObject(config).getJSONArray("plugins");
    String currentService = ApplicationContextHolder.getServiceName();
    for (plugin in plugins) {
        if (!plugin.targetServices.contains(currentService)) continue;
        if (!needUpdate(plugin.pluginId, plugin.checksum)) continue;
        File jar = minioService.download(plugin.jarUrl);
        pluginManager.unloadPlugin(plugin.pluginId);
        pluginManager.loadPlugin(jar);
    }
}''')

doc.add_page_break()

# ========== 第十二章 ==========
add_heading('第十二章  调试方案', 1)

add_heading('12.1 本地开发调试（首选）', 2)
add_para('插件加@Component + dev-mode=true，Spring注入，IDE正常debug：')
add_code('''@Component
@ConditionalOnProperty(name = "lowcode.plugin.dev-mode", havingValue = "true")
public class StockDeductExecutor implements ActionExecutor {
    // IDE里打断点，debug启动平台，正常F5/F6调试
}

// application.yml
lowcode:
  plugin:
    dev-mode: true   # true=本地debug，false=生产JAR加载''')

add_heading('12.2 四种调试方式', 2)
add_table(['方案', '场景', '能打断点', '推荐'], [
    ['本地@Component', '本地开发', '是', '首选'],
    ['远程JPDA', '生产排查', '是', '疑难问题'],
    ['日志SLF4J', '生产监控', '否', '生产首选'],
    ['设计器调试器', '联调验证', '否', '快速验证'],
])

doc.add_page_break()

# ========== 第十三章 ==========
add_heading('第十三章  风险与对策', 1)

risks = [
    ['SQL注入', '高', 'custom用参数化占位符；SqlFieldValidator校验表/列'],
    ['越权字段', '高', 'FieldGuard模型声明白名单过滤'],
    ['全表误删/误改', '高', 'delete/update强制conditions非空'],
    ['批量DoS', '中', 'max_batch_size限制(默认1000)，pageSize上限200'],
    ['批量转换性能', '高', 'BatchFieldTransformer批量预查询+内存映射'],
    ['触发环', '高', 'ActionContext.triggerChain检测重复中止'],
    ['事务内SQL过多', '高', '限制同步动作数量(建议≤5)；监控事务时间'],
    ['跨数据源无法同事务', '高', '同一触发链写操作必须同一数据源；跨源走异步'],
    ['脚本绕过事务', '高', '事务模式下PlatformApi只暴露compileXxx'],
    ['ClassLoader内存泄漏', '高', '卸载时关闭URLClassLoader；插件实现destroy()'],
    ['依赖冲突', '高', '独立ClassLoader隔离；公共依赖provided scope'],
    ['恶意插件', '高', '先不限制但全量审计；后续加JAR签名校验'],
    ['Nacos通知延迟', '中', '可接受最终一致；紧急手动重启'],
    ['MinIO下载失败', '中', '重试3次；失败告警；启动时补加载'],
    ['Groovy Metaspace OOM', '中', '脚本缓存按MD5去重；设上限LRU淘汰'],
    ['审批递归触发环', '高', 'triggerChain防环检测'],
]
add_table(['风险', '等级', '对策'], risks)

doc.add_page_break()

# ========== 第十四章 ==========
add_heading('第十四章  依赖清单', 1)

add_heading('14.1 插件项目依赖', 2)
add_table(['JAR', 'scope', '说明'], [
    ['plugin-api-1.0.jar', 'provided', '平台接口定义（ActionExecutor/PlatformApi/ActionContext）'],
    ['fastjson2-2.x.jar', 'provided', 'JSON处理'],
    ['slf4j-api.jar', 'provided', '日志（可选）'],
])

add_heading('14.2 平台新增依赖', 2)
add_table(['依赖', '版本', '大小', '用途'], [
    ['org.pf4j:pf4j', '3.9.0', '~200KB', '插件框架'],
    ['org.codehaus.groovy:groovy-all', '2.4.21', '~7MB', '脚本引擎'],
    ['com.googlecode.aviator:aviator', '5.3.3', '~500KB', '表达式引擎'],
])

doc.add_page_break()

# ========== 第十五章 ==========
add_heading('第十五章  落地步骤', 1)

steps = [
    ['1', 'DDL：全部表+ALTER', 'lowcode_tables.sql'],
    ['2', '配置+依赖引入', 'application.yml/pom.xml'],
    ['3', 'PlatformApi+缓存服务', 'engine/ 新增3文件'],
    ['4', '安全守卫：FieldGuard+SqlTemplateGuard', 'engine/ 新增2文件'],
    ['5', '模型管理+API元数据扩展', 'model/+registry/'],
    ['6', '执行引擎+字段转换', 'engine/ 新增4文件'],
    ['7', '对外入口OpenApiController', 'engine/'],
    ['8', '动作接口+8种执行器', 'trigger/action/'],
    ['9', '触发核心+TriggerApiController', 'trigger/'],
    ['10', 'MetaEngineService嵌入状态检测', 'engine/ 改'],
    ['11', '审批流', 'approval/ 新增4文件'],
    ['12', '插件管理+微服务分发', 'plugin/ 新增3文件'],
    ['13', '前端：全部设计器页面', 'web-code'],
    ['14', '验证：全链路联调', '-'],
]
add_table(['阶段', '内容', '文件'], steps)

doc.add_page_break()

# ========== 第十六章 FAQ ==========
add_heading('第十六章  FAQ', 1)

faqs = [
    ('Q: 一个按钮可以配多个插件吗？', 'A: 可以。actions_json是数组，可配任意多个动作，多个插件混着配，按顺序执行，同步动作在同一事务里。'),
    ('Q: 插件之间怎么传数据？', 'A: 通过ctx.variables。前一个动作put结果，后一个动作get取用。'),
    ('Q: 按钮点击后端会查数据库吗？', 'A: 不会。前端直接传当前行数据，后端直接用，不查库。动作里需要查其他数据由动作自己api.query()查。'),
    ('Q: 插件能直接连数据库吗？', 'A: 不能。必须通过PlatformApi的query/compileInsert/compileUpdate/compileDelete访问。'),
    ('Q: 插件能用Spring注解吗？', 'A: 生产环境不能（独立ClassLoader不在Spring容器）。本地开发可以加@Component调试。'),
    ('Q: 同步和异步动作区别？', 'A: 同步在事务内执行(compileXxx编译加入事务)，异步在事务提交后通过RocketMQ执行。'),
    ('Q: 插件更新需要重启吗？', 'A: 不需要。上传新JAR→Nacos广播→各服务自动热加载。'),
    ('Q: 脚本和插件怎么选？', 'A: 简单逻辑用Groovy脚本(前端编辑)，复杂业务用插件JAR(独立项目)。'),
    ('Q: 怎么本地调试插件？', 'A: dev-mode=true + @Component，IDE debug启动平台，正常打断点调试。'),
    ('Q: 事务里SQL太多怎么办？', 'A: 限制同步动作数量(建议≤5)。大量写操作拆分为多个触发或异步批量。'),
]
for q, a in faqs:
    add_para(q, bold=True)
    add_para(a)
    add_para('')

# ========== 总结 ==========
doc.add_page_break()
add_heading('总结', 1)
add_para('本文档涵盖低代码平台完整设计方案，核心要点：')
points = [
    '1. 元数据驱动CRUD：配置数据模型+API元数据→自动生成对外接口，零代码',
    '2. 通用触发源：按钮点击/状态变更/字段变更/API调用，4种触发类型',
    '3. 多动作+多插件：一个按钮可配任意多个动作，多个插件混配，按顺序执行',
    '4. 动作间传数据：ctx.variables传递，单据转换结果给后续插件用',
    '5. 事务化：写操作compileXxx编译收集，executeInTransaction原子提交',
    '6. 按钮直接传数据：前端传当前行数据，后端不查库',
    '7. 三层扩展：表达式(简单)→脚本(中等)→插件(复杂)，平台代码不动',
    '8. 插件通过PlatformApi访问平台能力，不耦合Spring，SPI自动发现',
    '9. 微服务分发：MinIO存储+Nacos广播+各服务监听热加载',
    '10. 审批流从零搭建：线性审批MVP，审批通过/拒绝触发新状态变更',
    '11. 缓存全部复用：L1Caffeine+L2Redis+BatchQueryStrategy，不造轮子',
]
for p in points:
    add_para(p)

output = r'D:\Documents\qwen-agent\code\code-service\lowcode-platform-plan\低代码平台完整开发文档.docx'
doc.save(output)
print(f'文档已生成：{output}')